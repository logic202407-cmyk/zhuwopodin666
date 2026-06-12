from pathlib import Path
import requests
import fitz
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FIG = ROOT / 'figures'
FIG.mkdir(exist_ok=True)
OUT = ROOT / '校园电动车充电安全监测与自动断电设计_课程设计报告_真实截图终稿.docx'

URLS = {
'front':'https://api.pdfmaker.to/download/51787e75-f4d3-46a2-b6d9-9396e8974b5f',
'back':'https://api.pdfmaker.to/download/43420cab-875c-4c40-9681-3b72f9b50a5c',
'basic_total':'https://api.pdfmaker.to/download/553102af-b2b0-4141-a28b-b5c31a4e158d',
'basic_current':'https://api.pdfmaker.to/download/c7fb2e1e-3c16-4115-a5be-3c7be690f973',
'ov_local':'https://api.pdfmaker.to/download/92d6c5ff-95e4-4da5-8a79-14cdc12b53af',
'oc_local':'https://api.pdfmaker.to/download/bd97121b-d037-435d-9bb7-3e24fcb9a2dd',
'ot_local':'https://api.pdfmaker.to/download/da200381-1d98-4085-863b-9e9f7f8bec3d',
'param_normal':'https://api.pdfmaker.to/download/440cceb7-61b4-4f77-a559-7261e7648ac9',
'param_critical':'https://api.pdfmaker.to/download/907ab58e-0b62-4d0e-887c-a7ee8b19b61b',
'param_current':'https://api.pdfmaker.to/download/09778e07-d468-4abd-8e86-02e5bc5df107',
'temp_plot':'https://api.pdfmaker.to/download/e0ef25e9-46f6-4295-8821-248404bfe0e9',
'temp_settings':'https://api.pdfmaker.to/download/81248027-0817-4ad3-967c-41006a2e5e3f',
'temp_table':'https://api.pdfmaker.to/download/54dfaef2-3d67-4198-ae26-9e7f6004acd4',
'worst':'https://api.pdfmaker.to/download/2b174b88-9864-4706-8631-7f3e41ed6fec',
'mc':'https://api.pdfmaker.to/download/ccac1152-b1d8-4220-b984-cf559fdbe95e',
'smoke_initial':'https://api.pdfmaker.to/download/7ceab7d0-b036-4cf5-bfc3-6cfe3860a2e9',
'smoke_final':'https://api.pdfmaker.to/download/58b27a93-572e-4552-bac9-a6efeed62f92'}

def get_image(name,url):
    pdf_path=FIG/f'{name}.pdf'; png_path=FIG/f'{name}.png'
    data=requests.get(url,timeout=90).content
    pdf_path.write_bytes(data)
    pdf=fitz.open(pdf_path); page=pdf[0]
    page.get_pixmap(matrix=fitz.Matrix(2.0,2.0),alpha=False).save(png_path)
    pdf.close(); return png_path

IMG={k:get_image(k,v) for k,v in URLS.items()}

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(3); sec.right_margin=Cm(2.5)
settings=doc.settings._element
u=OxmlElement('w:updateFields'); u.set(qn('w:val'),'true'); settings.append(u)

def font(run,size=12,bold=False,cn='宋体'):
    run.font.name='Times New Roman'; run._element.rPr.rFonts.set(qn('w:eastAsia'),cn); run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor(0,0,0)

def body(text,indent=True,size=12,align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p=doc.add_paragraph(); p.alignment=align; p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(0)
    if indent:p.paragraph_format.first_line_indent=Pt(24)
    font(p.add_run(text),size); return p

def heading(text,level):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_before=Pt(10 if level==1 else 6); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text); font(r,16 if level==1 else 14 if level==2 else 12,True,'黑体'); return p

def noborder(table):
    pr=table._tbl.tblPr; b=OxmlElement('w:tblBorders'); pr.append(b)
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'nil'); b.append(x)

def cell(c,text,bold=False):
    c.text=''; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); font(p.add_run(str(text)),10.5,bold)

def table3(title,headers,rows):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True; font(p.add_run(title),10.5)
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):cell(t.rows[0].cells[i],h,True)
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):cell(cs[i],v)
    pr=t._tbl.tblPr; b=OxmlElement('w:tblBorders'); pr.append(b)
    for e in ['top','bottom']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'10'); b.append(x)
    for e in ['left','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'nil'); b.append(x)
    for c in t.rows[0].cells:
        pr2=c._tc.get_or_add_tcPr(); tb=OxmlElement('w:tcBorders'); bt=OxmlElement('w:bottom'); bt.set(qn('w:val'),'single'); bt.set(qn('w:sz'),'6'); tb.append(bt); pr2.append(tb)
    doc.add_paragraph()

def eq(text,num):
    t=doc.add_table(rows=1,cols=2); noborder(t); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(text),12)
    p=t.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(p.add_run(num),12)

def fig(name,cap,maxw=15.7,maxh=18):
    path=IMG[name]
    with Image.open(path) as im:w,h=im.size
    width=maxw
    if width*h/w>maxh:width=maxh*w/h
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; noborder(t)
    trpr=t.rows[0]._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trpr.append(cant)
    p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1); p.add_run().add_picture(str(path),width=Cm(width))
    cp=t.cell(0,0).add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(6); font(cp.add_run(cap),10.5)

def page_num(p):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); b=OxmlElement('w:instrText'); b.text=' PAGE '; c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'end'); r._r.extend([a,b,c])

normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); normal.font.size=Pt(12); normal.paragraph_format.line_spacing=1.5
for n,s in [(1,16),(2,14),(3,12)]:
    st=doc.styles[f'Heading {n}']; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); st.font.size=Pt(s); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)

# cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(35); font(p.add_run('河南大学物理与电子学院'),20,True,'黑体')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25); font(p.add_run('《电子电路 CAD 技术》课程设计报告'),22,True,'黑体')
for _ in range(4):doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('校园电动车充电安全监测与自动断电设计'),18,True,'黑体')
for _ in range(5):doc.add_paragraph()
t=doc.add_table(rows=5,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; noborder(t)
for i,(a,b) in enumerate([('姓名','郑翔元'),('学号','________________'),('学院','物理与电子学院'),('专业','集成电路设计'),('提交日期','2026 年 6 月')]):
    t.cell(i,0).text=''; p=t.cell(i,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(p.add_run(a+'：'),14)
    t.cell(i,1).text=''; p=t.cell(i,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; font(p.add_run(b),14)
doc.add_page_break()

# abstract
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('摘  要'),16,True,'黑体')
body('针对校园电动车集中充电过程中可能出现的充电电压过高、充电电流过大和电池温度过高等安全隐患，设计了一种基于 OrCAD/PSpice 的多参数充电安全监测与自动断电电路。系统由前端检测、故障合成、充电允许控制、状态等效指示和负载断电五部分组成。前端采用 LM324 构成三路电压比较器，将 VBAT_SENSE、VI_SENSE 和 VTEMP 分别与 VOV、VOC、VOT 参考阈值比较，输出 OV、OC、OT 三路故障信号。后端使用 D1～D3 构成二极管或门，形成 FAULT_RAW，再由第四路 LM324 与 2.5 V 基准比较生成 EN。Q3 作为低端开关控制 220 Ω 等效负载；最终原理图中的 D4、D5 均为 1N4148，系统未使用蜂鸣器和继电器，自动断电由 Q3 截止实现。')
body('在标称瞬态仿真基础上，进一步完成三类故障局部动作分析、过压参考阈值参数扫描、−20～85 ℃温度扫描、最坏条件分析、100 次 Monte Carlo 容差分析及 Smoke 电应力分析。正常状态下 RLOAD 电流约为 22.42 mA；任一故障出现时，FAULT_RAW 升高、EN 降低，负载电流下降至微安量级。温度扫描中正常负载电流变化约 0.062%。Monte Carlo 结果中最大负载电流均值为 22.4201 mA，100 次运行均满足导通要求。初次 Smoke 分析发现 RLOAD、R4 和 R5 的功耗降额裕量不足；提高额定功率后，三者峰值功耗利用率分别降低至 83%、79% 和 79%，红色过应力项消失。结果表明，该电路能够实现三类故障联合检测和自动断电，并具有较好的温度稳定性、参数容差适应性和工程安全裕量。')
p=doc.add_paragraph(); font(p.add_run('关键词：'),12,True); font(p.add_run('OrCAD/PSpice；电动车充电；多参数监测；自动断电；可靠性分析'),12)
doc.add_page_break()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('目  录'),16,True,'黑体')
p=doc.add_paragraph(); r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); b=OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text=' TOC \\o "1-3" \\h \\z \\u '; c=OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'separate'); x=OxmlElement('w:t'); x.text='打开后右键更新整个目录'; d=OxmlElement('w:fldChar'); d.set(qn('w:fldCharType'),'end'); r._r.extend([a,b,c,x,d])

main=doc.add_section(WD_SECTION_START.NEW_PAGE); main.page_width=Cm(21); main.page_height=Cm(29.7); main.top_margin=Cm(2.5); main.bottom_margin=Cm(2.5); main.left_margin=Cm(3); main.right_margin=Cm(2.5); main.header.is_linked_to_previous=False; main.footer.is_linked_to_previous=False
pg=OxmlElement('w:pgNumType'); pg.set(qn('w:start'),'1'); main._sectPr.append(pg); hp=main.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(hp.add_run('河南大学物理与电子学院《电子电路 CAD 技术》课程设计报告'),10.5); page_num(main.footer.paragraphs[0])

heading('1 设计的目的',1); heading('1.1 设计背景',2)
body('随着校园电动车数量增加，宿舍区、教学区和生活区周边的集中充电需求不断增长。充电装置往往需要长时间连续运行，并可能受到电池老化、充电器异常、负载突变、环境温度升高、散热条件恶化以及传感器参数漂移等因素影响。当充电端电压超过允许范围、充电电流持续过大或电池温度异常升高时，若系统不能及时停止充电，可能造成电池性能下降、器件过应力，甚至形成安全隐患。因此，充电系统除完成能量传输外，还应具备状态采样、阈值判定、故障合成和执行切断等保护功能[1-3]。')
body('电子电路 CAD 技术能够在实际制作前完成电路结构设计、元件参数计算、网络连接检查、波形分析以及温度、容差和电应力验证。OrCAD Capture 适合进行模块化原理图绘制，PSpice 可执行瞬态、温度、参数扫描、Monte Carlo、Worst Case 和 Smoke 等分析[4-8]。本课题将模拟比较器、二极管逻辑和晶体管开关电路综合应用，形成一个具有明确工程功能和完整验证过程的课程设计。')
heading('1.2 设计目的',2)
body('本设计旨在完成一个能够识别过压、过流和过温故障，并在故障出现后自动切断等效负载的充电安全保护电路。通过设计掌握 Capture 中元件放置、网络命名、PARAM 参数元件、跨页连接、元件编号和设计规则检查；理解 LM324 比较器、二极管或门及 Q2N2222 低端开关的工作原理；能够根据负载电流、支路电流和功耗完成参数计算；能够建立多种仿真配置，并依据真实波形判断功能、稳定性和器件安全裕量。')
heading('1.3 设计意义',2)
body('本设计不是对真实高功率充电主回路的直接缩比，而是对安全监测和控制逻辑的等效验证。采用 VPULSE 模拟传感器输出，可将电压、电流和温度统一表示为比较器可处理的电压量，使研究重点集中在故障判定、逻辑合成和断电执行。三参数联合保护相对于单一保护具有更完整的故障覆盖范围；加入高温、容差、最坏条件和电应力分析，可使结论从标称功能正确扩展到参数变化时仍具可靠性和改进依据。')

heading('2 设计要求',1); heading('2.1 基本功能要求',2)
req=['建立 OV、OC、OT 三路故障检测通道。','过压输入 3.5 V/4.5 V，阈值 VOV=4.0 V。','过流输入 1.5 V/3.5 V，阈值 VOC=2.5 V。','过温输入 2.0 V/4.0 V，阈值 VOT=3.0 V。','任一路故障均应使 FAULT_RAW 变高。','正常时 EN 高、Q3 导通；故障时 EN 低、Q3 截止。','RLOAD=220 Ω，5 V 供电时正常电流约为 22 mA。']
for i,s in enumerate(req,1):body(f'（{i}）{s}',False)
heading('2.2 提高功能要求',2)
body('除标称瞬态仿真外，在 −20 ℃、27 ℃、40 ℃、60 ℃和 85 ℃条件下进行温度扫描；对 VOV 进行正常范围和临界范围参数扫描；综合高温、低供电、阈值偏高和负载偏差建立最坏条件测试；对关键电阻设置容差并执行 100 次 Monte Carlo；利用 Smoke 检查器件降额裕量，并根据初次结果完成参数修改和再次验证。')
heading('2.3 研究方案与技术路线',2)
body('研究按照“需求分析—功能分解—方案比较—参数计算—分模块绘图—工程合并—标称验证—环境与容差分析—电应力优化”的路线展开。三路 LM324 完成阈值比较，D1～D3 合成为 FAULT_RAW，U4A 生成 EN，Q1～Q3 完成状态支路与负载控制。标称功能正确后，再依次进行局部故障、参数、温度、最坏条件、Monte Carlo 和 Smoke 分析。')
heading('2.4 方案比较与特点',2)
body('检测环节可采用单片机采集后软件判断，也可采用模拟比较器直接判定。单片机方案便于显示和通信，但需要模数转换和程序设计；模拟比较器方案结构直观、响应路径短，且与课程内容契合，因此选择 LM324。故障合成可使用数字或门或二极管或门，后者器件少、结构简单，后级 2.5 V 阈值又可容忍正向压降，因此采用 1N4148 二极管或门。实际系统可使用继电器或功率 MOSFET，本仿真使用 Q2N2222 控制小电流等效负载。')

heading('3 设计内容',1); heading('3.1 电路组成及工作原理',2)
body('系统由前端检测、故障合成、充电允许控制、状态等效支路和负载断电模块组成。传感器等效信号接 LM324 同相端，参考电压接反相端，输入超过阈值时相应故障输出变高。D1、D2、D3 将三路故障汇接至 FAULT_RAW，R1=10 kΩ 在无故障时下拉该节点。')
eq('FAULT_RAW = OV ∨ OC ∨ OT','(3-1)'); eq('EN = ¬FAULT_RAW','(3-2)')
body('正常时三路故障输出均低，FAULT_RAW 被拉低，U4A 输出 EN 高电平，Q2 和 Q3 导通；任一故障出现时，FAULT_RAW 升高，EN 变低，Q2、Q3 截止，Q1 导通，Q3 截止后切断 RLOAD 电流。')
heading('3.2 前端检测模块',2)
body('过压通道 V2 由 3.5 V 跳变至 4.5 V，TD=2 ms、PW=2 ms，VOV=4 V；过流通道 V4 由 1.5 V 跳变至 3.5 V，TD=6 ms、PW=2 ms，VOC=2.5 V；过温通道 V6 由 2.0 V 跳变至 4.0 V，TD=10 ms、PW=2 ms，VOT=3 V。三路 LM324 均由 12 V 单电源供电。')
table3('表 3.1 三路检测通道参数',['通道','正常输入/V','故障输入/V','阈值/V','故障区间/ms'],[['过压','3.5','4.5','4.0','2～4'],['过流','1.5','3.5','2.5','6～8'],['过温','2.0','4.0','3.0','10～12']])
fig('front','图 3.1 前端过压、过流和过温检测模块原理图')
heading('3.3 故障合成与自动断电模块',2)
body('D1～D3 均采用 1N4148，阳极分别连接 OV、OC、OT，阴极汇接 FAULT_RAW。U4A 同相端接 2.5 V，反相端接 FAULT_RAW，正常时输出 EN 高电平，故障时输出低电平。Q1 基极经 R3=10 kΩ 接 FAULT_RAW；Q2、Q3 基极分别经 R6=10 kΩ、RB=10 kΩ 接 EN。R4、R5 均为 330 Ω；D4、D5 为普通 1N4148，不是 LED。最终电路没有继电器和蜂鸣器。')
fig('back','图 3.2 故障合成、状态等效支路与自动断电模块原理图')
heading('3.4 主要参数计算',2)
body('忽略 Q3 饱和压降时，等效负载的理论电流为：'); eq('I_LOAD = 5 / 220 = 22.73 mA','(3-3)')
body('仿真正常电流约为 22.42 mA，按该电流计算负载功耗为：'); eq('P_RLOAD = (22.42 mA)² × 220 Ω ≈ 0.1106 W','(3-4)')
body('按 D4/D5 正向压降 0.7 V、晶体管饱和压降 0.2 V 估算，330 Ω 支路电流约为 12.4 mA，电阻功耗约 0.051 W。10 kΩ 基极电阻在 LM324 高电平约 11 V 时可提供约 1.03 mA 基极电流，能够使 Q3 可靠导通。')
heading('3.5 OrCAD 工程实现与协同过程',2)
body('工程采用 PAGE1 和 PAGE2 两页结构。PAGE1 放置三路比较器及 OV、OC、OT 跨页端口，PAGE2 放置二极管或门、U4A、Q1～Q3 和 RLOAD。两人分别完成前端检测和后端控制，随后以同名 Off-Page Connector 合并。工程融合时删除后端测试用临时激励源，避免同一网络由多个理想源驱动，并执行 Annotate、DRC 和网表生成。设计中还解决了工程路径无法写入网表、打开工程连接服务器卡顿、探针显示和 Monte Carlo Measurement 设置等问题。')

heading('3.6 基本瞬态功能验证',2)
body('瞬态终止时间为 15 ms。0～2 ms 正常，2～4 ms 过压，4～6 ms 正常，6～8 ms 过流，8～10 ms 正常，10～12 ms 过温，12 ms 后恢复正常。')
fig('basic_total','图 3.3 基本瞬态仿真总波形')
body('真实波形显示，任一故障出现时 FAULT_RAW 升高至约 10.67 V，EN 降至毫伏量级，RLOAD 节点电压和负载电流同步降至近似零；故障消失后 EN 恢复至约 11.3 V，系统重新导通。')
fig('basic_current','图 3.4 等效负载电流瞬态波形')
body('正常状态 I(RLOAD) 约为 22.4 mA，三个故障窗口均降至近似零。边沿处少量离散采样点来自有限时间步长，不影响稳态判断。')
heading('3.7 三类故障局部动作分析',2)
fig('ov_local','图 3.5 过压故障局部保护波形'); body('2 ms 时过压通道翻转，FAULT_RAW 升高、EN 降低，4 ms 时恢复，动作区间与 V2 的延时和脉宽一致。')
fig('oc_local','图 3.6 过流故障局部保护波形'); body('6～8 ms 过流期间，OC 保持高电平，FAULT_RAW 与 EN 呈反相关系，负载被可靠切断。')
fig('ot_local','图 3.7 过温故障局部保护波形'); body('10～12 ms 过温期间，OT 触发同样的断电过程，证明任一路故障均可独立触发保护。')
heading('3.8 参数扫描分析',2)
body('将 VOV 设为全局参数。第一组按 3.8～4.2 V、步长 0.1 V 扫描，因各阈值均低于 4.5 V 故障输入，全部曲线均可触发保护。')
fig('param_normal','图 3.8 过压阈值正常范围参数扫描结果')
body('第二组按 4.3～4.7 V 扫描。VOV 小于 4.5 V 时能够动作，高于 4.5 V 时过压输入不再超过阈值，2～4 ms 内负载不能切断；VOV=4.5 V 位于临界比较区。')
fig('param_critical','图 3.9 过压阈值临界范围参数扫描结果'); fig('param_current','图 3.10 临界阈值扫描下的负载电流结果')
body('参数扫描说明，阈值过高会造成漏保护，实际设计必须为传感器误差、参考源偏差和器件失调预留足够裕量。')
heading('3.9 温度扫描与高温稳定性',2)
body('温度扫描依次取 −20 ℃、27 ℃、40 ℃、60 ℃和 85 ℃，重复执行同一瞬态分析。')
fig('temp_settings','图 3.11 温度扫描仿真设置'); fig('temp_plot','图 3.12 −20～85 ℃温度扫描总波形'); fig('temp_table','图 3.13 温度扫描测量数据')
body('不同温度曲线基本重合。正常 I(RLOAD) 约由 22.427 mA 下降至 22.413 mA，变化量 0.014 mA，相对变化约 0.062%；EN 高电平由约 11.372 V 下降至 11.241 V，仍具有足够驱动裕量；FAULT_RAW 无故障低电平远低于 2.5 V 阈值。需要指出，理想 VPULSE 不包含真实传感器温漂，因此硬件还需考虑基准源、传感器和 PCB 散热。')
heading('3.10 最坏条件分析',2)
body('最坏条件设置为 85 ℃、VCC12=10.8 V、VCC5=4.75 V、VOV=4.2 V、VOC=2.7 V、VOT=3.2 V、RLOAD=242 Ω，三路故障高电平分别为 4.3 V、2.8 V 和 3.3 V，仅比阈值高 0.1 V。正常理论电流为 4.75/242=19.63 mA。')
fig('worst','图 3.14 最坏条件下系统保护波形')
body('三段故障期间 FAULT_RAW 仍能升高、EN 仍能降至接近零，负载电流被切断，未出现漏动作，说明所设最坏条件下电路仍保持保护能力。')
heading('3.11 Monte Carlo 容差分析',2)
body('对关键电阻设置 5%～10% 容差并进行 100 次随机运行，测量 Max/Min(V(EN))、Max/Min(V(FAULT_RAW)) 和 Max/Min(I(RLOAD))。')
fig('mc','图 3.15 最大负载电流的 Monte Carlo 统计分布')
body('Max(I(RLOAD)) 范围为 21.2549～23.6988 mA，均值 22.4201 mA，标准差约 0.6036 mA，Yield 为 100%；Min(I(RLOAD)) 约为 −1.8533～−1.6490 μA，负号仅表示参考方向。FAULT_RAW 高电平均值约 10.6749 V，低电平均值约 0.641 mV；EN 高电平均值约 11.3149 V，低电平均值约 5.7324 mV。随机误差只造成小范围分散，没有破坏保护逻辑。')
heading('3.12 Smoke 电应力分析与改进',2)
body('“Smoke Analysis succeeded”仅表示分析运行成功，必须结合颜色和 % Max 判断器件安全。初次结果中 RLOAD、R4 和 R5 的 PDM 项出现红色过应力。')
fig('smoke_initial','图 3.16 初始额定功率设置下的 Smoke 分析结果')
body('RLOAD 峰值功耗约 110.5824 mW，R4、R5 峰值功耗约 53.87 mW，但在既定降额规则下安全裕量不足。因此将 RLOAD 的 POWER 由 0.25 W 提高至 0.5 W，将 R4、R5 由 0.125 W 提高至 0.25 W，MAX-TEMP 保持 125 ℃。')
fig('smoke_final','图 3.17 提高额定功率后的 Smoke 分析结果')
body('重新分析后红色过应力项消失。RLOAD 峰值功耗利用率约 83%，R4、R5 均约 79%，说明提高额定功率后获得合理裕量。部分运放项目仍显示黄色 100%，输出窗口提示 RCA、RJC 热参数缺失，因此本次结论主要针对已配置额定参数的电阻功耗，不能替代真实器件数据手册和 PCB 热设计。')
heading('3.13 设计特点、创新性与局限',2)
body('本设计将过压、过流和过温统一为电压比较问题，以二极管或门实现联合触发；采用两页模块化设计和 PARAM 统一管理阈值、电源与负载；在基本瞬态之外完成临界阈值、全温度、最坏条件、100 次 Monte Carlo 以及 Smoke 改进，形成从功能到稳定性、容差和安全裕量的多层次验证闭环。')
body('局限在于传感器仍由理想源代替，比较器未加入迟滞，故障消失后自动恢复而未锁存，D4/D5 不是 LED，Q3 只能驱动小电流等效负载。后续可加入真实分压、电流采样、NTC、滤波、施密特迟滞、SR 锁存、隔离驱动及功率 MOSFET。')

heading('4 总结与感悟',1); heading('4.1 设计结论',2)
body('本课程设计完成了校园电动车充电安全监测与自动断电电路的方案比较、原理图设计、参数计算和多层次仿真。真实 OrCAD 原理图与 PSpice 波形表明，三类故障均能被正确识别，任一故障均可使 FAULT_RAW 变高、EN 变低并切断 RLOAD。温度扫描中负载电流变化约 0.062%，最坏条件下未出现漏动作，100 次 Monte Carlo 结果保持稳定。Smoke 初次分析发现功耗降额问题，提高额定功率后红色过应力项消失。由此可认为该设计达到了多参数监测和自动断电的预期目标。')
heading('4.2 课程目标达成与个人感悟',2)
body('通过本次设计，掌握了 Capture 工程管理、跨页端口、PARAM 元件、属性设置、网表检查以及 PSpice 瞬态、参数、温度、Worst Case、Monte Carlo 和 Smoke 分析。工程合并、网络重复驱动、网表生成、探针显示和高级分析参数等问题的排查说明，电子电路 CAD 不只是画图，而是需要使电路原理、模型、参数、仿真条件与结论相互一致。前后端分工和最终联合验证也锻炼了模块化设计、协作排查及技术文档撰写能力。')
body('后续可将理想采样源替换为真实传感器电路，增加输入滤波、迟滞和故障锁存；功率执行部分应使用隔离继电器或功率 MOSFET，并依据数据手册、环境温度和 PCB 散热完成降额设计。')

heading('参考文献',1)
refs=['[1] 康华光，陈大钦，张林. 电子技术基础：模拟部分[M]. 北京：高等教育出版社，2021.','[2] 阎石. 数字电子技术基础[M]. 北京：高等教育出版社，2016.','[3] 王兆安，刘进军. 电力电子技术[M]. 北京：机械工业出版社，2009.','[4] 邱关源，罗先觉. 电路[M]. 北京：高等教育出版社，2022.','[5] Cadence Design Systems. PSpice User Guide[EB/OL]. San Jose: Cadence Design Systems, 2024.','[6] 河南大学物理与电子学院. 电子电路 CAD 技术课程讲义：时域分析[Z]. 开封：河南大学，2026.','[7] 河南大学物理与电子学院. 电子电路 CAD 技术课程讲义：温度、参数与统计分析[Z]. 开封：河南大学，2026.','[8] 河南大学物理与电子学院. 电子电路 CAD 技术课程讲义：高级分析[Z]. 开封：河南大学，2026.']
for x in refs:
    p=body(x,False,10.5); p.paragraph_format.hanging_indent=Pt(21)

doc.core_properties.title='校园电动车充电安全监测与自动断电设计'; doc.core_properties.author='郑翔元'; doc.core_properties.subject='电子电路 CAD 技术课程设计报告（真实截图版）'
doc.save(OUT)
print(OUT)
